from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

slides_content = [
    {
        "title": "FinRisk AI: Unified NLP Risk Engine & Portfolio Stress-Testing Platform",
        "subtitle": "S&P Global & CRISIL Campus Hackathon 2026 | Candidate: Harshita Singla | Indira Gandhi Delhi Technical University for Women",
        "bullets": [
            "Track: AI/NLP Risk Engine with Strategic Portfolio Stress Testing (Module B)",
            "Transforms unstructured market feeds into quantitative credit risk and regulatory capital triggers.",
            "Combines Hugging Face FinBERT domain sentiment with Basel III Expected Credit Loss (ECL) mechanics."
        ]
    },
    {
        "title": "Problem Statement & The Latency Gap in Risk Monitoring",
        "subtitle": "Why Traditional Credit Reviews Lag Behind Real-World Market Distress",
        "bullets": [
            "Unstructured Data Lag: Critical credit events appear in breaking news and social feeds 48-72 hours before rating agency adjustments.",
            "Information Asymmetry: Banks manage large wholesale portfolios without real-time linkage between qualitative news and portfolio capital adequacy.",
            "Our Solution: An automated, sub-second NLP Risk Engine mapping directly to wholesale asset repricing and capital adequacy (CRAR) buffers."
        ]
    },
    {
        "title": "System Architecture & High-Performance Data Flow",
        "subtitle": "End-to-End Pipeline: Ingestion -> NLP Engine -> Structured JSON -> Portfolio Stress Engine",
        "bullets": [
            "Data Ingestion: Dual-channel streaming ingesting financial news and micro-blog feeds into a normalized schema.",
            "Risk NLP Engine: FinBERT sentiment scoring (S in [-1, 1]), zero-shot event taxonomy classification, and explainable impact scoring.",
            "Stress-Testing Engine: Automated shock propagation executing Basel III ECL on loans and Modified Duration repricing on debt.",
            "Consumption Layer: Interactive Streamlit analytics dashboard and production-grade FastAPI REST endpoints."
        ]
    },
    {
        "title": "AI/NLP Methodology & Impact Scoring Mechanics",
        "subtitle": "Explainable Quantitative Scoring Replacing Black-Box Heuristics",
        "bullets": [
            "Sentiment Analysis: FinBERT model outputs calibrated score S = P(pos) - P(neg), delivering domain-accurate scores (e.g. -0.954).",
            "Event Categorization: Predefined taxonomy spanning Credit Events, Macroeconomic rate shifts, Geopolitical sanctions, and Earnings.",
            "Explainable Impact Score (1-10): I = Clip[1,10](0.50 * BaseSeverity + 0.30 * (|S| * 10) + 0.20 * SourceWeight).",
            "Attribution & Evidence: Engine automatically extracts driver tokens (e.g. 'debt default', 'spread widening') for risk committee auditability."
        ]
    },
    {
        "title": "Wholesale Banking Portfolio & Basel III Stress Simulation",
        "subtitle": "Institutional-Grade Risk Modeling Over Arbitrary Percentage Cuts",
        "bullets": [
            "Synthetic Portfolio Exposure: ₹100 Lakhs distributed across Corporate Loans, G-Secs, Corporate Bonds, Equities, and Swaps.",
            "Basel III Credit Impairment: dLoss = EAD * dPD * LGD (applied at 45% standard institutional LGD).",
            "Modified Duration Repricing: dP = -D_mod * dy * P, accurately modeling bond price erosion under interest rate hikes.",
            "Regulatory Capital Impact: Evaluates Tier-1 Capital depletion and automatically flags breaches against the 11.5% minimum CRAR requirement."
        ]
    },
    {
        "title": "Prototype Demonstration & Key Evaluation Results",
        "subtitle": "Empirical Verification Across Diverse Market Scenarios",
        "bullets": [
            "Credit Event Shock: Headline regarding commercial defaults triggered 9/10 impact, reducing portfolio value from ₹100L to ₹90.59L (-9.4%).",
            "Capital Buffer Alert: Stressed CRAR declined from 17.65% to 6.35%, automatically triggering an institutional regulatory capital alert.",
            "Macroeconomic Rate Shock: 75 bps surprise rate hike dynamically repriced 10Y Sovereign G-Secs based on duration (7.2y) without touching unrelated credit loans.",
            "False-Positive Resilience: Product launches and positive earnings retain 'NOMINAL' status without executing capital stress."
        ]
    },
    {
        "title": "Domain Impact, Governance & Production Roadmap",
        "subtitle": "Scaling from Prototype to Enterprise Risk Management (ERM)",
        "bullets": [
            "Business Value: Provides chief risk officers (CROs) with dynamic, intraday scenario stress-testing rather than static quarterly reviews.",
            "Enterprise Integration: Built-in FastAPI endpoints ready for streaming enterprise Kafka message queues.",
            "Governance & Regulatory Compliance: Full audit trail of extracted tokens, confidence scores, and formula breakdowns ready for regulatory examination.",
            "Future Roadmap: Multi-entity relationship graphs, Basel IV output floors, and real-time counterparty exposure limits."
        ]
    }
]

for item in slides_content:
    slide = prs.slides.add_slide(blank_layout)
    
    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = item["title"]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)
    
    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = item["subtitle"]
    p2.font.size = Pt(15)
    p2.font.color.rgb = RGBColor(37, 99, 235)
    
    # Content Box
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5))
    ctf = content_box.text_frame
    ctf.word_wrap = True
    for bullet in item["bullets"]:
        bp = ctf.add_paragraph()
        bp.text = f"• {bullet}"
        bp.font.size = Pt(17)
        bp.space_after = Pt(16)
        bp.font.color.rgb = RGBColor(51, 65, 85)

prs.save("docs/presentation.pptx")
print("✅ Saved docs/presentation.pptx successfully!")